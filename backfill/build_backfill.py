"""Builds the MVP's starting data from the Financial Statement spreadsheet.

    python3 -I backfill/build_backfill.py "Untitled spreadsheet.xlsx" "Oct 26" [--embed]

Reads one monthly tracking tab (default "Oct 26") and "RAB Bulanan", applies the
categorisation agreed with the user, checks that every pocket ends on the sheet's own
closing balance, and writes backfill/<tab>.json plus the copy embedded in mvp/index.html.

Rules (from the user, 2026-10-28):
- Every RAB AR item on a date splits that day's matching transaction: the AR amount goes to
  the AR category (linked to the AR item), the rest is the user's own, in Lifestyle.
- "Reimburse" rows are AR settlements.
- Kos income/expense get their own "Kos" category; Groceries (200.000/month) and Lifestyle
  (coffee, for himself) are new; the 10–12 Oct holiday purchases count as Date.
- AP "Titip" items have no transaction ("Not sure").
"""
import datetime as dt
import json
import re
import sys
import uuid
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
ARGS = [a for a in sys.argv[1:] if not a.startswith('--')]
XLSX = Path(ARGS[0]) if ARGS else ROOT / 'Untitled spreadsheet.xlsx'
TAB = ARGS[1] if len(ARGS) > 1 else 'Oct 26'
if not XLSX.is_absolute():
    XLSX = ROOT / XLSX

POCKETS = ['Kos', 'BCA', 'Mandiri', 'Cash', 'Gopay', 'ShopeePay', 'Ovo']
CREDIT_COLS = 'DEFGHIJ'   # money in, per pocket, in POCKETS order
DEBIT_COLS = 'KLMNOPQ'    # money out
CLOSING_COLS = 'TUVWXYZ'  # running balance per pocket

EXPENSE = ['Office lunch', 'Office transport', 'Gas/toll', 'Date', 'Hobby', 'Spotify', 'Gym', 'Claude',
           'Groceries', 'Lifestyle', 'Health', 'Admin', 'Kos', 'Others']
INCOME = ['Salary', 'Kos', 'Others']

# RAB "Rincian Keperluan" → app category, with the RAB's price × qty per unit.
BUDGET = [('Office lunch', 30000, 'Day', 22), ('Office transport', 20000, 'Day', 22), ('Gas/toll', 1000000, 'Month', 1),
          ('Date', 1000000, 'Month', 1), ('Hobby', 500000, 'Month', 1), ('Spotify', 79900, 'Month', 1),
          ('Gym', 185000, 'Month', 1), ('Claude', 360000, 'Month', 1), ('Groceries', 200000, 'Month', 1),
          ('Health', 0, 'Month', 1), ('Admin', 0, 'Month', 1)]

BY_NAME = {
    'makan siang': 'Office lunch', 'sarapan': 'Office lunch', 'makan siang + jajan': 'Office lunch',
    'bensin': 'Gas/toll', 'etoll': 'Gas/toll', 'emoney': 'Gas/toll', 'parkir': 'Gas/toll',
    'parkir + charger': 'Gas/toll', 'cuci mobil': 'Gas/toll', 'motor': 'Gas/toll',
    'makan malam': 'Date', 'hotel': 'Date', 'parma': 'Date', "gia's taman": 'Date', "gia's coffee": 'Date',
    'nasi pedas': 'Date', 'sinar warisan': 'Date', 'blue doors': 'Date', 'mimadre': 'Date',
    'solaria + holland + minum': 'Date', 'wingstop': 'Lifestyle', 'tennis & makan': 'Date', 'alfamart (exclude)': 'Date',
    'bola': 'Hobby', 'tennis': 'Hobby', 'jajan tennis': 'Hobby', 'restring': 'Hobby',
    'kopi': 'Lifestyle', 'spotify': 'Spotify', 'claude': 'Claude',
    'galon': 'Groceries', 'air': 'Groceries', 'wifi': 'Kos',
    # Lifestyle: spending on himself — snacks and personal care
    'nipple cover': 'Lifestyle', 'carb + permen karet': 'Lifestyle', 'norrm': 'Lifestyle', 'potong': 'Lifestyle',
    'masker': 'Lifestyle', 'liq': 'Lifestyle', 'liquid': 'Lifestyle', 'parfum': 'Lifestyle',
    'rombak celana': 'Lifestyle', 'sendal': 'Lifestyle',
    'obat': 'Health', 'admin': 'Admin',
}
INCOME_BY_NAME = {'gaji': 'Salary', 'kos': 'Kos'}

# Per-row corrections from the category review (row number in the tab → category).
ROW_CATEGORY = {
    112: ('expense', 'Lifestyle'), 115: ('expense', 'Lifestyle'),  # his own lunches, not office lunch
}

# RAB AR items matched to the tracking row they came out of (row number in the tab).
# The AR amounts in RAB are the source of truth, even where the transaction differs (row 99).
AR_ROWS = {
    'Makan siang mama': 5, 'Laundry@2026-10-02': 10, 'Regal': 36, 'Laundry@2026-10-09': 41,
    'Pinjem Akbar': 82, 'Telor@2026-10-18': 86, 'Sarapan + makan siang': 89, 'Makan malam@2026-10-21': 99,
    'Makan siang': 100, 'Makan malam@2026-10-24': 111, 'Makan malam@2026-10-25': 113, 'Telor@2026-10-26': 116,
    'Indomie + shampoo': 117, 'Bedak': 120, 'Lampu': 123, 'Bensin tol': 119,
}


def new_id():
    return str(uuid.uuid4())


def money(v):
    return round(float(v or 0), 2)


wb = openpyxl.load_workbook(XLSX, data_only=True)
ws, rab = wb[TAB], wb['RAB Bulanan']

# ---- read the tracking tab -------------------------------------------------------------
rows, date = {}, None
for r in range(3, ws.max_row + 1):
    name = ws[f'C{r}'].value
    if ws[f'B{r}'].value:
        date = ws[f'B{r}'].value.date().isoformat()
    deltas = {}
    for i, p in enumerate(POCKETS):
        d = money(ws[f'{CREDIT_COLS[i]}{r}'].value) - money(ws[f'{DEBIT_COLS[i]}{r}'].value)
        if round(d, 2):
            deltas[p] = round(d, 2)
    if name and deltas:
        rows[r] = {'row': r, 'date': date, 'name': str(name).strip(), 'deltas': deltas}
last_row = max(r for r in range(3, ws.max_row + 1) if ws[f'T{r}'].value is not None)
closing = {p: money(ws[f'{CLOSING_COLS[i]}{last_row}'].value) for i, p in enumerate(POCKETS)}

# ---- read RAB AR / AP -------------------------------------------------------------------
def rab_list(first, last):
    out = []
    for r in range(first, last + 1):
        q, name, amount = rab[f'Q{r}'].value, rab[f'R{r}'].value, rab[f'S{r}'].value
        if q and name and amount:
            out.append({'date': q.date().isoformat(), 'name': str(name).strip(), 'amount': money(amount)})
    return out

ar_items, ap_items = rab_list(7, 26), rab_list(31, 45)
income_budget = money(rab['G16'].value)
target_eomb = money(rab['N19'].value)

# ---- model -----------------------------------------------------------------------------
pockets = {}
for i, p in enumerate(POCKETS):
    pockets[p] = {'id': new_id(), 'name': p, 'parent_id': None, 'sort': i, 'archived': False}
pockets['Kos']['parent_id'] = pockets['BCA']['id']
# BCA first, then its Kos sub-pocket, as on the Balance screen.
pockets['BCA']['sort'], pockets['Kos']['sort'] = 0, 1

cats = {}
for i, n in enumerate(EXPENSE):
    cats[('expense', n)] = {'id': new_id(), 'name': n, 'kind': 'expense', 'sort': 999 if n == 'Others' else i, 'archived': False}
for i, n in enumerate(INCOME):
    cats[('income', n)] = {'id': new_id(), 'name': n, 'kind': 'income', 'sort': 999 if n == 'Others' else 100 + i, 'archived': False}
for i, n in enumerate(['AR', 'AP']):
    cats[('debt', n)] = {'id': new_id(), 'name': n, 'kind': 'debt', 'sort': 1000 + i, 'archived': False}

transactions, debts, seq = [], [], [0]


def add_tx(date, name, cat, lines):
    """lines: {pocket: signed amount}. cat: (kind, name) or None. Returns the transaction."""
    seq[0] += 1
    t = {'id': new_id(), 'date': date, 'name': name, 'category_id': cats[cat]['id'] if cat else None,
         'created_at': f'{date}T08:00:00.{seq[0]:03d}Z',
         'lines': [{'id': new_id(), 'pocket_id': pockets[p]['id'], 'amount': round(a, 2)} for p, a in lines.items() if round(a, 2)]}
    transactions.append(t)
    return t


def transfer(date, src, dst, amount, fee=0):
    return add_tx(date, f'Top up from {src} to {dst}', None, {src: -(amount + fee), dst: amount})


def dinner_cat(amount):
    """A dinner under 100.000 is the user's own (Lifestyle); a bigger one is a date."""
    return 'Lifestyle' if amount < 100000 else 'Date'


def default_cat(row):
    if row['row'] in ROW_CATEGORY:
        return ROW_CATEGORY[row['row']]
    name = row['name'].lower()
    if all(v > 0 for v in row['deltas'].values()):
        return ('income', INCOME_BY_NAME.get(name, 'Others'))
    if name == 'makan malam':
        return ('expense', dinner_cat(-sum(row['deltas'].values())))
    if name == 'bensin':  # small fills are the work motorbike, bigger ones the car
        return ('expense', 'Office transport' if -sum(row['deltas'].values()) < 60000 else 'Gas/toll')
    if name == 'makan':  # a weekday meal is office lunch; at the weekend it follows the dinner rule
        weekday = dt.date.fromisoformat(row['date']).weekday() < 5
        return ('expense', 'Office lunch' if weekday else dinner_cat(-sum(row['deltas'].values())))
    if name in BY_NAME:
        return ('expense', BY_NAME[name])
    return ('expense', 'Lifestyle' if name.startswith('jajan') else 'Others')


def ar_debt(item, open_tx=None, settled=None, settle_tx=None):
    debts.append({'id': new_id(), 'kind': 'AR', 'date': item['date'], 'name': item['name'], 'amount': item['amount'],
                  'settled_at': settled, 'open_tx_id': open_tx['id'] if open_tx else None,
                  'settle_tx_id': settle_tx['id'] if settle_tx else None, 'created_at': f"{item['date']}T00:00:00Z"})


# AR items keyed the way AR_ROWS names them.
ar_by_key, seen = {}, {}
for item in ar_items:
    dupes = [i for i in ar_items if i['name'] == item['name']]
    key = item['name'] if len(dupes) == 1 else f"{item['name']}@{item['date']}"
    ar_by_key[key] = item
assert set(ar_by_key) == set(AR_ROWS), set(ar_by_key) ^ set(AR_ROWS)
ar_for_row = {row: ar_by_key[k] for k, row in AR_ROWS.items() if row}

opening = rows.pop(min(rows))
assert opening['name'].lower() == 'saldo awal', opening
add_tx(opening['date'], 'Opening balance', None, opening['deltas'])


def handle(row):
    r, d, n, dl = row['row'], row['date'], row['name'], row['deltas']
    # Mixed rows: split into what actually happened.
    if r == 4:    # Kos pays BCA back, plus a BCA charge
        transfer(d, 'Kos', 'BCA', dl['BCA'] + 6000)
        add_tx(d, 'Admin', ('expense', 'Admin'), {'BCA': -6000})
    elif n.lower() == 'reimburse':  # settles an AR from before this month
        t = add_tx(d, f'Settled: {n}', ('debt', 'AR'), dl)
        ar_debt({'date': d, 'name': n, 'amount': round(sum(dl.values()), 2)}, settled=d, settle_tx=t)
    elif r == 23:  # card spend with a tiny credit back on the same row
        add_tx(d, n, ('expense', 'Lifestyle'), {'BCA': dl['BCA']})
    elif r == 53:  # top up ShopeePay, then dinner from it
        transfer(d, 'Mandiri', 'ShopeePay', 100000, 1000)
        add_tx(d, n, ('expense', dinner_cat(54000)), {'ShopeePay': -54000})
    elif r in (54, 77):  # cash handed over, paid back by transfer
        transfer(d, 'Cash', 'BCA', dl['BCA'])
    elif r == 91:
        transfer(d, 'Mandiri', 'ShopeePay', 200000, 1000)
    elif r == 95:
        transfer(d, 'Mandiri', 'Cash', 300000)
    elif r == 99:  # dinner for mum; RAB records the AR as 32.000, the sheet shows 27.000 leaving ShopeePay
        t = add_tx(d, n, ('debt', 'AR'), {'ShopeePay': -27000})
        ar_debt(ar_for_row[99], open_tx=t)
        add_tx(d, n, ('income', 'Others'), {'BCA': 6000})
    elif r == 107:
        transfer(d, 'Mandiri', 'Gopay', 200000, 1200)
        add_tx(d, n, ('expense', 'Others'), {'Gopay': -60500})
    elif r == 120:  # handled with 121: the 201.000 was a ShopeePay top-up for mum's bedak
        pass
    elif r == 121:
        transfer(d, 'Mandiri', 'ShopeePay', 200000)
        bedak = ar_for_row[120]
        t = add_tx(d, rows[120]['name'], ('debt', 'AR'), {'ShopeePay': -169400, 'Mandiri': -1000})
        assert abs(169400 + 1000 - bedak['amount']) < 0.01
        ar_debt(bedak, open_tx=t)
        transfer(d, 'BCA', 'Cash', 200000)
        add_tx(d, 'Jajan', ('expense', 'Lifestyle'), {'BCA': -21900})
    elif r == 129:
        transfer(d, 'Mandiri', 'Gopay', 300000, 1200)
        add_tx(d, n, ('expense', 'Claude'), {'Gopay': -349000})
    elif r in ar_for_row:  # AR part to AR, the rest is the user's own (Lifestyle)
        item = ar_for_row[r]
        (pocket, amount), = dl.items()
        assert -amount >= item['amount'] - 0.01, (r, item)
        assert item['date'] == d, (r, item)
        t = add_tx(d, item['name'], ('debt', 'AR'), {pocket: -item['amount']})
        ar_debt(item, open_tx=t)
        rest = round(-amount - item['amount'], 2)
        if rest:
            add_tx(d, n, ('expense', 'Lifestyle'), {pocket: -rest})
    else:
        signs = {v > 0 for v in dl.values()}
        assert len(signs) == 1, f'row {r} mixes money in and out and has no rule: {row}'
        add_tx(d, n, default_cat(row), dl)


for r in sorted(rows):
    handle(rows[r])

for item in ap_items:
    debts.append({'id': new_id(), 'kind': 'AP', 'date': item['date'], 'name': item['name'], 'amount': item['amount'],
                  'settled_at': None, 'open_tx_id': None, 'settle_tx_id': None, 'created_at': f"{item['date']}T00:00:00Z"})

# ---- checks ----------------------------------------------------------------------------
bal = {p: 0.0 for p in POCKETS}
pid = {v['id']: k for k, v in pockets.items()}
for t in transactions:
    for l in t['lines']:
        bal[pid[l['pocket_id']]] += l['amount']
problems = {p: (round(bal[p], 2), closing[p]) for p in POCKETS if abs(bal[p] - closing[p]) > 0.01}
assert not problems, f'closing balances differ from the sheet (app, sheet): {problems}'
open_ar = sum(d['amount'] for d in debts if d['kind'] == 'AR' and not d['settled_at'])
dropped = sum(ar_by_key[k]['amount'] for k, row in AR_ROWS.items() if row is None)
rab_ar = money(rab['S27'].value)
assert abs(open_ar + dropped - rab_ar) < 0.01, (open_ar, dropped, rab_ar)

month = transactions[0]['date'][:8] + '01'
period = {'id': new_id(), 'month': month, 'income': income_budget, 'target_eomb': target_eomb}
items = [{'id': new_id(), 'period_id': period['id'], 'category_id': cats[('expense', n)]['id'],
          'price': price, 'unit': unit, 'qty': qty, 'sort': i} for i, (n, price, unit, qty) in enumerate(BUDGET)]

out = {'v': 4, 'source': f'{XLSX.name} · {TAB}', 'pockets': list(pockets.values()), 'categories': list(cats.values()),
       'transactions': transactions, 'debts': debts, 'periods': [period], 'items': items}
for coll in ('pockets', 'categories'):
    for x in out[coll]:
        x['created_at'] = f'{month}T00:00:00Z'

slug = re.sub(r'\W+', '-', TAB.lower()).strip('-')
dest = ROOT / 'backfill' / f'{slug}.json'
dest.write_text(json.dumps(out, ensure_ascii=False, indent=1))

# Embed into the prototype between its markers, only once the categories are signed off.
if '--embed' in sys.argv:
    html_path = ROOT / 'mvp' / 'index.html'
    html = html_path.read_text()
    start, end = '/*BACKFILL*/', '/*END BACKFILL*/'
    a, b = html.index(start) + len(start), html.index(end)
    html_path.write_text(html[:a] + json.dumps(out, ensure_ascii=False, separators=(',', ':')) + html[b:])

kinds = {}
for t in transactions:
    c = next((c for c in cats.values() if c['id'] == t['category_id']), None)
    k = c['name'] + ('' if c['kind'] != 'income' else ' (in)') if c else 'Transfer/opening'
    kinds[k] = kinds.get(k, 0) + 1
print(f'{len(transactions)} transactions, {len(debts)} debts → {dest.relative_to(ROOT)}' + (' and mvp/index.html' if '--embed' in sys.argv else ''))
print('closing balances match the sheet:', {p: round(bal[p], 2) for p in POCKETS})
print(f'open AR {open_ar:,.0f} = RAB AR total {rab_ar:,.0f} minus {dropped:,.0f} that was not AR')
print('per category:', dict(sorted(kinds.items(), key=lambda x: -x[1])))
